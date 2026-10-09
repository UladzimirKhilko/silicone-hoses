"""Паспорт и бирки к отгрузке (docs/10) в формате Word — как образцы заказчика.

Правила (решения 07.10.2026):
- масса на бирке — брутто коробки из упаковочного листа завода; неполная коробка —
  брутто ÷ шт. в коробке × количество;
- номер паспорта — атрибут партии; если в отгрузке несколько партий — по паспорту на каждую.
"""

import io
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .models import CartonSpec

LOGO = Path(__file__).parent / "assets" / "bsi-logo.png"
COMPANY = "ОАО «БПА БЕЛСТРОЙИНДУСТРИЯ»"
COMPANY_CONTACTS = (
    "220075, Республика Беларусь, г.Минск, ул.Селицкого, 17А т/ф.:(+375 17) 364 82 24; 364 76 82\n"
    "E-mail: bsi@bpabsi.by   http: www.bpabsi.by"
)
PRODUCT_NAME = "Патрубок силиконовый"


# --- Раскладка по коробкам ---

@dataclass
class Box:
    product: object
    batch: object
    qty: int
    gross_kg: Decimal | None
    count: int  # сколько таких коробок (бирок)


def main_spec(specs):
    """Основная коробка — та, которых в партии больше всего; при равенстве — самая вместительная."""
    return max(specs, key=lambda s: (s.cartons or 0, s.qty_per_carton)) if specs else None


def _round_kg(value):
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def box_plan(lines):
    """Строки отгрузки → коробки. Одинаковые коробки одного изделия и партии объединяются."""
    specs = {}
    for s in CartonSpec.objects.filter(
        batch_id__in={line.batch_id for line in lines}, product_id__in={line.product_id for line in lines}
    ):
        specs.setdefault((s.batch_id, s.product_id), []).append(s)
    boxes = []
    for line in lines:
        spec = main_spec(specs.get((line.batch_id, line.product_id), []))
        if spec is None:
            boxes.append(Box(line.product, line.batch, line.qty, None, 1))
            continue
        full, rest = divmod(line.qty, spec.qty_per_carton)
        if full:
            boxes.append(Box(line.product, line.batch, spec.qty_per_carton, spec.gross_kg, full))
        if rest:
            weight = _round_kg(spec.gross_kg / spec.qty_per_carton * rest)
            boxes.append(Box(line.product, line.batch, rest, weight, 1))
    return boxes


def format_kg(value):
    """16.00 → «16,0»; 5.65 → «5,65»; 21.40 → «21,4» — как на бирках заказчика."""
    if value is None:
        return "____"
    text = f"{_round_kg(value):.2f}".rstrip("0")
    if text.endswith("."):
        text += "0"
    return text.replace(".", ",")


def product_code(product):
    return product.code if product.code.upper().startswith("BSI") else f"BSI {product.code}"


# --- Word: общие мелочи ---

def _new_document(left_cm, right_cm, top_cm, bottom_cm, size=13):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(left_cm), Cm(right_cm)
    section.top_margin, section.bottom_margin = Cm(top_cm), Cm(bottom_cm)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(size)
    style.paragraph_format.space_after = Pt(0)
    return doc


def _cell_text(cell, text, bold=False, size=None, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    for i, part in enumerate(text.split("\n")):
        if i:
            paragraph.add_run().add_break()
        run = paragraph.add_run(part)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
    if align is not None:
        paragraph.alignment = align
    return paragraph


def _no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _set_widths(table, widths):
    """Ширины колонок и в сетке таблицы (её читает LibreOffice), и в ячейках (её читает Word)."""
    table.autofit = False
    for grid_col, width in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        grid_col.set(qn("w:w"), str(int(width.emu / 635)))  # EMU → twips
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def _rule(doc):
    """Горизонтальная линия под шапкой — нижняя граница абзаца."""
    p = doc.add_paragraph()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", "000000")):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    p._p.get_or_add_pPr().append(borders)
    return p


def _to_bytes(doc):
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- Паспорт ---

def _passport_page(doc, shipment, batch, lines):
    head = doc.add_table(rows=1, cols=2)
    _no_borders(head)
    _set_widths(head, (Cm(2.6), Cm(15.1)))
    head.cell(0, 1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if LOGO.exists():
        head.cell(0, 0).paragraphs[0].add_run().add_picture(str(LOGO), width=Cm(2.0))
    _cell_text(head.cell(0, 1), COMPANY, bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    contacts = head.cell(0, 1).add_paragraph()
    contacts.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, part in enumerate(COMPANY_CONTACTS.split("\n")):
        if i:
            contacts.add_run().add_break()
        r = contacts.add_run(part)
        r.bold, r.font.size = True, Pt(10)
    _rule(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(14)
    title.paragraph_format.space_after = Pt(14)
    r = title.add_run("Паспорт")
    r.bold, r.font.size = True, Pt(16)

    info = doc.add_table(rows=3, cols=1)
    info.style = "Table Grid"
    _set_widths(info, (Cm(17.5),))
    number = batch.passport_number or "________"
    _cell_text(info.cell(0, 0), f"№ {number} от {shipment.date:%d.%m.%Y}г", bold=True)
    recipient = shipment.customer.full_name or shipment.customer.name
    p = _cell_text(info.cell(1, 0), "Получатель продукции: ")
    p.add_run(recipient).bold = True
    _cell_text(info.cell(2, 0), f"Отгрузочный документ:  ТТН   № {shipment.ttn_number or ''}")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    items = doc.add_table(rows=1 + len(lines), cols=4)
    items.style = "Table Grid"
    items.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (Cm(8.4), Cm(2.4), Cm(3.8), Cm(2.8))
    for i, title_text in enumerate(("Наименование продукции:", "Кол-во штук:", "Дата изготовления:", "Номер партии:")):
        _cell_text(items.cell(0, i), title_text, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row, ln in enumerate(lines, start=1):
        _cell_text(items.cell(row, 0), f"{PRODUCT_NAME}\n{product_code(ln.product)}")
        _cell_text(items.cell(row, 1), str(ln.qty), align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_text(items.cell(row, 2), batch.manufactured or "", align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_text(items.cell(row, 3), batch.number or "", align=WD_ALIGN_PARAGRAPH.CENTER)
    if len(lines) > 1:  # одна партия на паспорт — дата и номер общие, как в образце
        for col in (2, 3):
            merged = items.cell(1, col).merge(items.cell(len(lines), col))
            _cell_text(merged, (batch.manufactured if col == 2 else batch.number) or "", align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_widths(items, widths)
    for row in items.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for text in ("Детали соответствуют требованиям: КД", "Отпуск разрешил:"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(28)
        r = p.add_run(text)
        r.bold, r.font.size = True, Pt(14)
    sign = doc.add_table(rows=2, cols=3)
    _no_borders(sign)
    _set_widths(sign, (Cm(5.8), Cm(5.8), Cm(5.8)))
    for i, caption in enumerate(("должность", "подпись", "ФИО")):
        _cell_text(sign.cell(0, i), "\n" + "_" * 18, align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_text(sign.cell(1, i), caption, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)


def passport_docx(shipment):
    """Паспорт к отгрузке. Несколько партий — несколько паспортов в одном файле, каждый на своей странице."""
    lines = list(shipment.lines.select_related("product", "batch").order_by("batch__received_date", "batch_id", "pk"))
    by_batch = {}
    for ln in lines:
        by_batch.setdefault(ln.batch, []).append(ln)
    doc = _new_document(2.0, 1.25, 1.25, 1.25)
    for i, (batch, batch_lines) in enumerate(by_batch.items()):
        if i:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _passport_page(doc, shipment, batch, batch_lines)
    return _to_bytes(doc)


# --- Бирки ---

def label_text(box):
    batch = box.batch
    return [
        ("Наименование: ", f"{PRODUCT_NAME}", True),
        ("", f"{product_code(box.product)}", True),
        ("Номер партии: ", batch.number or "____", False),
        ("Количество: ", f"{box.qty} шт.", False),
        ("Масса: ", f"{format_kg(box.gross_kg)} кг", False),
        ("Дата изготовления: ", f"{batch.manufactured}г." if batch.manufactured else "____", False),
        ("ТНПА: ", "КД", False),
        ("Контроллер: ", "", False),
    ]


def labels_docx(shipment):
    """Бирки: слева текст бирки, справа — сколько таких бирок (коробок), как в образце заказчика."""
    lines = list(shipment.lines.select_related("product", "batch").order_by("product__code", "batch_id", "pk"))
    boxes = box_plan(lines)
    doc = _new_document(3.0, 1.5, 2.0, 2.0, size=12)
    doc.styles["Normal"].font.name = "Arial"
    table = doc.add_table(rows=len(boxes), cols=2)
    table.style = "Table Grid"
    for row, box in zip(table.rows, boxes):
        cell = row.cells[0]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        for i, (label, value, bold) in enumerate(label_text(box)):
            if i:
                paragraph = cell.add_paragraph()
            if i == 1:
                paragraph.paragraph_format.left_indent = Cm(2.6)
            paragraph.add_run(label)
            paragraph.add_run(value).bold = bold
        cell.add_paragraph()
        _cell_text(row.cells[1], str(box.count), align=WD_ALIGN_PARAGRAPH.CENTER)
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_widths(table, (Cm(12.7), Cm(3.8)))
    return _to_bytes(doc)

