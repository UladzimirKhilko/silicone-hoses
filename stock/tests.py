import datetime as dt
import io
from decimal import Decimal

import openpyxl
import pytest
from django.core.management import call_command
from django.urls import reverse

from catalog.models import Customer, Product

from . import services
from .models import Adjustment, Batch, ReceiptLine, Shipment, ShipmentLine


@pytest.fixture
def mtz(db):
    return Customer.objects.create(name="МТЗ")


@pytest.fixture
def product(mtz):
    return Product.objects.create(code="BSI D75L100-2", yapib="24.1119", customer=mtz)


@pytest.fixture
def client_in(client, django_user_model):
    client.force_login(django_user_model.objects.create_user("m", password="pass12345"))
    return client


def received(product, qty, date, **kw):
    batch = Batch.objects.create(status=Batch.Status.RECEIVED, received_date=date, **kw)
    ReceiptLine.objects.create(batch=batch, product=product, qty_expected=qty, qty_received=qty)
    return batch


def test_pick_prefers_single_batch_over_old_leftovers():
    old, new = object(), object()
    assert services.pick_batches([(old, 3), (new, 950)], 200) == [(new, 200)]
    assert services.pick_batches([(old, 3), (new, 950)], 2) == [(old, 2)]
    assert services.pick_batches([(old, 3), (new, 5)], 7) == [(old, 3), (new, 4)]
    assert services.pick_batches([(old, 3)], 7) is None


def test_balances_from_documents(product, mtz):
    samples = received(product, 4, dt.date(2025, 2, 8))
    serial = received(product, 950, dt.date(2025, 9, 18), number="02/25")
    in_way = Batch.objects.create()
    ReceiptLine.objects.create(batch=in_way, product=product, qty_expected=500)
    ship = Shipment.objects.create(customer=mtz, date=dt.date(2025, 11, 19))
    services.post_shipment(ship, [{"product": product, "qty": 200}])
    assert ship.lines.get().batch == serial
    adj = Adjustment.objects.create(kind=Adjustment.Kind.DEFECT, date=dt.date(2025, 11, 20))
    services.post_adjustment(adj, [{"product": product, "qty": 1, "batch": samples}])
    assert services.on_hand()[product.pk] == 4 + 950 - 200 - 1
    assert services.in_transit()[product.pk] == 500


def test_cannot_ship_more_than_stock(product, mtz):
    received(product, 10, dt.date(2025, 1, 1))
    ship = Shipment.objects.create(customer=mtz, date=dt.date(2025, 2, 1))
    with pytest.raises(services.InsufficientStock):
        services.post_shipment(ship, [{"product": product, "qty": 6}, {"product": product, "qty": 5}])


def test_explicit_batch_is_respected_and_checked(product, mtz):
    a = received(product, 10, dt.date(2025, 1, 1))
    b = received(product, 10, dt.date(2025, 2, 1))
    ship = Shipment.objects.create(customer=mtz, date=dt.date(2025, 3, 1))
    services.post_shipment(ship, [{"product": product, "qty": 8, "batch": b}, {"product": product, "qty": 10}])
    assert sorted((line.batch_id, line.qty) for line in ship.lines.all()) == sorted([(b.pk, 8), (a.pk, 10)])


def test_movements_running_balance(product, mtz):
    received(product, 100, dt.date(2025, 1, 1))
    ship = Shipment.objects.create(customer=mtz, date=dt.date(2025, 2, 1))
    services.post_shipment(ship, [{"product": product, "qty": 30}])
    assert [(m.qty, m.balance) for m in services.movements(product)] == [(100, 100), (-30, 70)]


# --- Экраны ---

def test_receive_with_shortage(client_in, product):
    batch = Batch.objects.create(invoice="TS1")
    line = ReceiptLine.objects.create(batch=batch, product=product, qty_expected=100)
    client_in.post(reverse("stock:batch_receive", args=[batch.pk]), {"received_date": "2025-09-18", f"received_{line.pk}": "97", f"landed_{line.pk}": "3,47"})
    batch.refresh_from_db()
    line.refresh_from_db()
    assert batch.is_received and line.qty_received == 97 and line.shortage == 3
    assert line.landed_price_byn == Decimal("3.47")
    assert services.on_hand()[product.pk] == 97


def test_create_batch_in_transit(client_in, product):
    data = {"invoice": "TS2", "factory_orders": "", "number": "", "manufactured": "", "passport_number": "", "notes": "",
            "lines-TOTAL_FORMS": "2", "lines-INITIAL_FORMS": "0", "lines-0-product": product.pk, "lines-0-qty_expected": "300",
            "lines-0-price_cny": "6.67", "lines-1-product": "", "lines-1-qty_expected": ""}
    resp = client_in.post(reverse("stock:batch_create"), data)
    assert resp.status_code == 302
    assert services.in_transit()[product.pk] == 300 and services.on_hand().get(product.pk, 0) == 0


def _shipment_post(client, customer, product, qty, price=""):
    return client.post(reverse("stock:shipment_create"), {
        "customer": customer.pk, "date": "2025-11-19", "ttn_number": "5134450", "notes": "",
        "lines-TOTAL_FORMS": "1", "lines-INITIAL_FORMS": "0",
        "lines-0-product": product.pk, "lines-0-qty": qty, "lines-0-price_byn": price, "lines-0-batch": "",
    })


def test_shipment_view_auto_batch_and_last_price(client_in, product, mtz):
    received(product, 500, dt.date(2025, 9, 18), number="02/25")
    _shipment_post(client_in, mtz, product, 50, "8.14")
    _shipment_post(client_in, mtz, product, 200)  # цена пустая — берётся прошлая
    lines = ShipmentLine.objects.order_by("pk")
    assert [line.price_byn for line in lines] == [Decimal("8.14"), Decimal("8.14")]
    assert services.on_hand()[product.pk] == 250


def test_shipment_view_refuses_overselling(client_in, product, mtz):
    received(product, 10, dt.date(2025, 9, 18))
    resp = _shipment_post(client_in, mtz, product, 11)
    assert resp.status_code == 200 and not Shipment.objects.exists()
    assert "Не хватает товара" in resp.content.decode()


def test_delete_shipment_returns_stock(client_in, product, mtz):
    received(product, 10, dt.date(2025, 9, 18))
    _shipment_post(client_in, mtz, product, 4)
    client_in.post(reverse("stock:shipment_delete", args=[Shipment.objects.get().pk]))
    assert services.on_hand()[product.pk] == 10


def test_free_samples_adjustment(client_in, product, mtz):
    received(product, 10, dt.date(2025, 9, 18))
    client_in.post(reverse("stock:adjustment_create"), {
        "kind": "free_sample", "date": "2025-10-01", "customer": mtz.pk, "notes": "на согласование",
        "lines-TOTAL_FORMS": "1", "lines-INITIAL_FORMS": "0", "lines-0-product": product.pk, "lines-0-qty": "3", "lines-0-batch": "",
    })
    assert services.on_hand()[product.pk] == 7


def test_list_shows_stock_and_exports_excel(client_in, product):
    received(product, 42, dt.date(2025, 9, 18))
    body = client_in.get(reverse("catalog:product_list")).content.decode()
    assert ">42<" in body
    resp = client_in.get(reverse("stock:stock_export"))
    ws = openpyxl.load_workbook(io.BytesIO(resp.content)).active
    assert [c.value for c in ws[2]][0] == "BSI D75L100-2" and ws["G2"].value == 42


def test_product_card_shows_movements(client_in, product, mtz):
    received(product, 100, dt.date(2025, 9, 18), number="02/25")
    resp = client_in.get(reverse("catalog:product_detail", args=[product.pk]))
    assert resp.context["on_hand"] == 100 and len(resp.context["movements"]) == 1


# --- Перенос из Excel ---

def test_import_mtz_excel(tmp_path, product, mtz):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B4"], ws["C4"], ws["L4"], ws["AJ4"] = "Патрубок силиконовый BSI D75L100-2", 6.67, 8.14, 3.47
    ws["D4"], ws["E4"], ws["F4"] = 3, 4, 950   # приходы: образцы, TS2407202, партия 02/25
    ws["M4"], ws["W4"], ws["X4"] = 4, 50, 200   # отгрузки 08.07, 11.11, 19.11
    path = tmp_path / "mtz.xlsx"
    wb.save(path)
    call_command("import_mtz_excel", str(path), stdout=io.StringIO())
    assert services.on_hand()[product.pk] == 3 + 4 + 950 - 4 - 50 - 200
    nov = Shipment.objects.get(date=dt.date(2025, 11, 19))
    assert nov.ttn_number == "5134450"
    assert [(line.batch.number, line.qty) for line in nov.lines.all()] == [("02/25", 200)]
    assert ReceiptLine.objects.get(batch__number="02/25").landed_price_byn == Decimal("3.47")


def test_ttn_scan_upload_and_view(client_in, product, mtz, settings, tmp_path):
    from django.core.files.uploadedfile import SimpleUploadedFile

    settings.MEDIA_ROOT = tmp_path
    ship = Shipment.objects.create(customer=mtz, date=dt.date(2025, 11, 11))
    client_in.post(reverse("stock:shipment_attach_ttn", args=[ship.pk]),
                   {"ttn_number": "5134413", "ttn_scan": SimpleUploadedFile("t.pdf", b"%PDF scan")})
    ship.refresh_from_db()
    assert ship.ttn_number == "5134413"
    resp = client_in.get(reverse("stock:shipment_ttn_scan", args=[ship.pk]))
    assert b"".join(resp.streaming_content) == b"%PDF scan"
    client_in.logout()
    assert client_in.get(reverse("stock:shipment_ttn_scan", args=[ship.pk])).status_code == 302


def test_lists_sorted_newest_first(client_in, product, mtz):
    received(product, 100, dt.date(2025, 1, 1))
    for day in (5, 20, 10):
        ship = Shipment.objects.create(customer=mtz, date=dt.date(2025, 3, day))
        services.post_shipment(ship, [{"product": product, "qty": 1}])
    dates = [s.date.day for s in client_in.get(reverse("stock:shipment_list")).context["shipments"]]
    assert dates == [20, 10, 5]


# --- Паспорт и бирки (этап 3) ---

@pytest.fixture
def batch_0225(product, mtz):
    mtz.full_name = "ОАО «Минский тракторный завод»"
    mtz.save()
    batch = received(product, 950, dt.date(2025, 9, 18), number="02/25", manufactured="07.2025", passport_number="09/002/25")
    from .models import CartonSpec

    CartonSpec.objects.create(batch=batch, product=product, qty_per_carton=85, gross_kg=Decimal("16.0"))
    return batch


def _ship(customer, product, qty, ttn="5134450", batch=None):
    ship = Shipment.objects.create(customer=customer, date=dt.date(2025, 11, 19), ttn_number=ttn)
    services.post_shipment(ship, [{"product": product, "qty": qty, "batch": batch}])
    return ship


def test_box_plan_reproduces_real_labels(batch_0225, product, mtz):
    from .paperwork import box_plan, format_kg

    boxes = box_plan(list(_ship(mtz, product, 200).lines.all()))
    assert [(b.qty, format_kg(b.gross_kg), b.count) for b in boxes] == [(85, "16,0", 2), (30, "5,65", 1)]
    boxes = box_plan(list(_ship(mtz, product, 50, ttn="5134413").lines.all()))
    assert [(b.qty, format_kg(b.gross_kg), b.count) for b in boxes] == [(50, "9,41", 1)]


def test_main_carton_is_the_most_common(batch_0225, product):
    from .models import CartonSpec
    from .paperwork import main_spec

    CartonSpec.objects.filter(batch=batch_0225).update(cartons=1)
    CartonSpec.objects.create(batch=batch_0225, product=product, qty_per_carton=40, gross_kg=Decimal("8"), cartons=10)
    assert main_spec(list(CartonSpec.objects.filter(batch=batch_0225))).qty_per_carton == 40


def test_box_plan_without_carton_data(product, mtz):
    from .paperwork import box_plan

    received(product, 100, dt.date(2025, 1, 1))
    boxes = box_plan(list(_ship(mtz, product, 30).lines.all()))
    assert [(b.qty, b.gross_kg, b.count) for b in boxes] == [(30, None, 1)]


def _docx_text(data):
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def test_passport_contents(batch_0225, product, mtz):
    from .paperwork import passport_docx

    text = _docx_text(passport_docx(_ship(mtz, product, 200)))
    for expected in ("№ 09/002/25 от 19.11.2025г", "ОАО «Минский тракторный завод»", "№ 5134450",
                     "BSI D75L100-2", "200", "07.2025", "02/25", "Детали соответствуют требованиям: КД"):
        assert expected in text


def test_one_passport_per_batch(batch_0225, product, mtz):
    from .paperwork import passport_docx

    old = received(product, 10, dt.date(2025, 2, 8), number="01/25", passport_number="09/001/25")
    ship = Shipment.objects.create(customer=mtz, date=dt.date(2025, 11, 19))
    services.post_shipment(ship, [{"product": product, "qty": 5, "batch": old}, {"product": product, "qty": 100, "batch": batch_0225}])
    text = _docx_text(passport_docx(ship))
    assert text.count("Паспорт") == 2 and "09/001/25" in text and "09/002/25" in text


def test_labels_contents(batch_0225, product, mtz):
    import docx

    from .paperwork import labels_docx

    d = docx.Document(io.BytesIO(labels_docx(_ship(mtz, product, 200))))
    rows = d.tables[0].rows
    assert len(rows) == 2
    assert "Количество: 85 шт." in rows[0].cells[0].text and "Масса: 16,0 кг" in rows[0].cells[0].text
    assert rows[0].cells[1].text == "2"
    assert "Количество: 30 шт." in rows[1].cells[0].text and "Масса: 5,65 кг" in rows[1].cells[0].text
    assert "Номер партии: 02/25" in rows[1].cells[0].text and "Дата изготовления: 07.2025г." in rows[1].cells[0].text


def test_paperwork_downloads_and_warnings(client_in, batch_0225, product, mtz):
    ship = _ship(mtz, product, 200, ttn="")
    resp = client_in.get(reverse("stock:shipment_detail", args=[ship.pk]))
    assert resp.context["box_total"] == 3
    assert any("ТТН" in w for w in resp.context["paperwork_warnings"])
    for name in ("stock:shipment_passport", "stock:shipment_labels"):
        r = client_in.get(reverse(name, args=[ship.pk]))
        assert r.status_code == 200 and r["Content-Type"].startswith("application/vnd.openxmlformats")
        assert b"".join(r.streaming_content)[:2] == b"PK"


def test_batch_info_and_carton_editing(client_in, product):
    from .models import CartonSpec

    batch = received(product, 100, dt.date(2025, 9, 18))
    client_in.post(reverse("stock:batch_info", args=[batch.pk]), {"number": "03/25", "manufactured": "10.2025", "passport_number": "09/003/25"})
    batch.refresh_from_db()
    assert (batch.number, batch.manufactured, batch.passport_number) == ("03/25", "10.2025", "09/003/25")
    url = reverse("stock:carton_add", args=[batch.pk])
    client_in.post(url, {"product": product.pk, "qty_per_carton": 85, "gross_kg": "16,0"})
    client_in.post(url, {"product": product.pk, "qty_per_carton": 85, "gross_kg": "17"})  # дубль — отклоняется
    assert list(CartonSpec.objects.values_list("qty_per_carton", "gross_kg")) == [(85, Decimal("16.00"))]


def test_decimal_comma_accepted_in_shipment_price(client_in, product, mtz):
    received(product, 100, dt.date(2025, 9, 18))
    _shipment_post(client_in, mtz, product, 10, "8,14")
    assert ShipmentLine.objects.get().price_byn == Decimal("8.14")
