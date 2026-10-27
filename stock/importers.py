"""Разбор файлов завода: инвойс (русский, xlsx) и упаковочный лист (xls/xlsx).

Форматы — по поставке TS2407202 (docs/08). Колонки ищутся по заголовкам, а не по номерам,
чтобы мелкие изменения шаблона завода не ломали загрузку.

Особенности упаковочного листа:
- одно изделие повторяется во многих строках (разные коробки) — суммируем;
- смешанные коробки (образцы): первая строка группы с маркировкой коробки, весом и общим
  количеством, следующие строки — только модель и количество в этой коробке;
- справа бывают служебные промежуточные итоги, внизу — строка «IN TOTAL» — пропускаем.
"""

import datetime as dt
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path


class ParseError(Exception):
    """Файл не похож на ожидаемый документ."""


# --- Чтение таблиц ---

def read_rows(file, name=""):
    """Первый лист .xlsx/.xls или таблицы PDF → список строк (списки значений).

    Тип файла определяется по содержимому, а не по имени: имя может быть любым.
    """
    data = file.read() if hasattr(file, "read") else Path(file).read_bytes()
    if data[:5] == b"%PDF-":
        return _pdf_rows(data)
    if data[:4] == b"\xd0\xcf\x11\xe0":  # старый формат Excel (.xls)
        import xlrd

        sheet = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
        return [sheet.row_values(r) for r in range(sheet.nrows)]
    import io

    import openpyxl

    if data[:2] != b"PK":
        raise ParseError("Файл не похож на Excel (.xlsx, .xls) или PDF.")
    ws = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True).worksheets[0]
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _pdf_rows(data):
    """Таблицы всех страниц PDF подряд (продолжение таблицы на следующей странице — тоже).

    Читается PDF с текстом (выгруженный из Excel/программы). Скан-картинку прочитать нельзя.
    """
    import pymupdf

    doc = pymupdf.open(stream=data, filetype="pdf")
    rows, text = [], ""
    for page in doc:
        page_text = page.get_text()
        if not text:  # шапка первой страницы (номер, дата) часто напечатана вне таблицы
            rows.extend([line] for line in page_text.splitlines()[:40] if line.strip())
        text += page_text
        for table in page.find_tables().tables:
            rows.extend(table.extract())
    if not text.strip():
        raise ParseError("PDF — это скан (картинка без текста), прочитать его нельзя. "
                         "Попросите у завода Excel или PDF, сохранённый из Excel.")
    if not any(len(r) > 2 for r in rows):
        raise ParseError("В PDF не нашлось таблицы. Нужен инвойс или упаковочный лист с таблицей позиций.")
    return rows


def _text(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""


def _num(v):
    if v in (None, ""):
        return None
    if isinstance(v, float):
        # Excel хранит числа как float: 2951.3999999999996 → 2951.4; цены завода неокруглённые
        # (12.6760942760943) — 10 знаков сохраняем.
        return Decimal(repr(round(v, 10))).normalize() if v != int(v) else Decimal(int(v))
    if isinstance(v, int):
        return Decimal(v)
    text = re.sub(r"[\s\u00a0]", "", str(v))
    text = re.sub(r"^(?:[¥￥$€]|CNY|RMB|USD|BYN)|(?:[¥￥$€]|CNY|RMB|USD|BYN)$", "", text, flags=re.IGNORECASE)
    if re.fullmatch(r"-?\d{1,3}(?:,\d{3})+(?:\.\d+)?", text):  # 1,334.95 и 1,500 — запятая-разделитель тысяч (документы завода)
        text = text.replace(",", "")
    elif not re.fullmatch(r"-?\d+(?:[.,]\d+)?", text):
        return None
    try:
        return Decimal(text.replace(",", "."))
    except InvalidOperation:
        return None


def _key(text):
    return re.sub(r"\s+", "", str(text).lower())


def _find_header(rows, must, limit=120, without=()):
    """Строка заголовков: в ней есть хотя бы одно слово из каждой группы `must` и нет слов `without`.
    Пробелы и переносы строк в заголовках не важны («QTY/\nCTN» = «qty/ctn»)."""
    for i, row in enumerate(rows[:limit]):
        cells = [_key(_text(c)) for c in row]
        if all(any(_key(w) in c for c in cells for w in (group if isinstance(group, tuple) else (group,))) for group in must) \
                and not any(_key(w) in c for c in cells for w in without):
            return i, cells
    return None, None


def _col(cells, *words, exclude=()):
    for i, c in enumerate(cells):
        if any(_key(w) in c for w in words) and not any(_key(x) in c for x in exclude):
            return i
    return None


# --- Инвойс ---

@dataclass
class InvoiceLine:
    name: str
    qty: int
    price: Decimal | None
    amount: Decimal | None
    hs_code: str = ""
    unit: str = ""  # «м/п» — количество в метрах (рукава БелАЗ), обычно пусто = штуки


@dataclass
class Invoice:
    number: str = ""
    date: dt.date | None = None
    contract: str = ""
    currency: str = ""
    lines: list = field(default_factory=list)

    @property
    def total(self):
        return sum((line.amount or 0) for line in self.lines)


RU_MONTHS = {"янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "мая": 5, "июн": 6, "июл": 7, "авг": 8, "сен": 9,
             "окт": 10, "ноя": 11, "дек": 12}


def _parse_date(text):
    text = text.strip()
    for fmt in ("%Y/%m/%d", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    # «июнь.19.2026», «сент.28.2026» — так дату пишет завод в русских документах
    if m := re.match(r"([а-яё]+)\.?\s*(\d{1,2})[.,\s]+(\d{4})", text.lower()):
        month = RU_MONTHS.get(m.group(1)[:3])
        if month:
            try:
                return dt.date(int(m.group(3)), month, int(m.group(2)))
            except ValueError:
                return None
    return None


# Колонка модели в китайском упаковочном листе: «matched models» или «Item No.» (другая версия шаблона)
MODEL_WORDS = ("model", "itemno", "item no")
NAME_WORDS = ("наименование", "description", "commodity", "goods", "品名", "名称")
QTY_WORDS = ("к-во", "к-в", "кол-во", "количество", "quantity", "qty", "数量")
PRICE_WORDS = ("цена", "unitprice", "price", "单价")
AMOUNT_WORDS = ("сумма", "amount", "totalvalue", "金额")


def parse_invoice(rows):
    head, cells = _find_header(rows, [NAME_WORDS, QTY_WORDS, PRICE_WORDS], without=("qty/ctn",))
    if head is None:
        raise ParseError("Не нашёл в файле таблицу с колонками «Наименование», «Кол-во», «Цена» — это точно инвойс?")
    c_name = _col(cells, *NAME_WORDS)
    c_qty = _col(cells, *QTY_WORDS)
    c_price = _col(cells, *PRICE_WORDS)
    c_amount = _col(cells, *AMOUNT_WORDS)
    c_hs = _col(cells, "тнвэд", "тн вэд", "hscode", "hs")
    c_unit = _col(cells, "ед.изм", "едизм", "unit", exclude=("price", "цена"))
    inv = Invoice()
    header_text = " ".join(_text(c) for row in rows[:head] for c in row)
    if m := re.search(r"(?:Инвойс|Invoice)\s*(?:№|No\.?|Nr\.?|#)?\s*[.:]?\s*([A-Za-z]{1,4}(?:-[A-Za-z]{1,4})*-?\d{5,}[A-Za-z0-9\-/]*)",
                      header_text, re.IGNORECASE):  # TS2407202, JH-J-20261009
        inv.number = m.group(1)
    elif m := re.search(r"\b([A-Z]{2}\d{6,})\b", header_text):
        inv.number = m.group(1)
    if m := re.search(r"(?:Дата|Date)\s*:?\s*([\d./\-]{8,10})", header_text, re.IGNORECASE):
        inv.date = _parse_date(m.group(1))
    if m := re.search(r"к договору\s*([\w/\-]+\s*от\s*[\d.]+)", header_text, re.IGNORECASE):
        inv.contract = m.group(1)
    if m := re.search(r"\((CNY|USD|EUR|RMB)\)", " ".join(cells), re.IGNORECASE):
        inv.currency = m.group(1).upper().replace("RMB", "CNY")
    elif re.search(r"[¥￥]", " ".join(_text(c) for row in rows[head:head + 5] for c in row)):
        inv.currency = "CNY"
    for row in rows[head + 1:]:
        name = _text(row[c_name]) if c_name is not None and c_name < len(row) else ""
        qty = _num(row[c_qty]) if c_qty is not None and c_qty < len(row) else None
        total_row = name.lower().startswith(("итого", "total", "всего"))
        if not name or qty is None or total_row:
            if inv.lines and total_row:
                break
            continue
        inv.lines.append(InvoiceLine(
            name=name,
            qty=int(qty),
            price=_num(row[c_price]) if c_price is not None else None,
            amount=_num(row[c_amount]) if c_amount is not None else None,
            hs_code=_text(row[c_hs]) if c_hs is not None else "",
            unit=_text(row[c_unit]) if c_unit is not None and c_unit < len(row) else "",
        ))
    if not inv.lines:
        raise ParseError("В инвойсе не нашлось строк с изделиями.")
    return inv


# --- Упаковочный лист завода ---

@dataclass
class PackingCarton:
    qty_per_carton: int
    cartons: int
    gross_kg: Decimal | None
    net_kg: Decimal | None


@dataclass
class PackingItem:
    name: str
    mark: str = ""  # маркировка на коробке (кит. описание)
    color: str = ""
    qty: int = 0
    cartons: list = field(default_factory=list)  # только коробки с одним изделием
    in_mixed: int = 0  # штук в смешанных коробках


@dataclass
class PackingList:
    title: str = ""
    invoice: str = ""
    items: dict = field(default_factory=dict)  # name -> PackingItem
    total_cartons: int | None = None
    total_qty: int | None = None
    total_gross: Decimal | None = None
    total_net: Decimal | None = None
    mixed_cartons: list = field(default_factory=list)  # [(маркировка, [(name, qty)])]
    source: str = "cn"  # cn — китайский лист (раскладка по коробкам), ru — русский (итоги по изделию)
    date: dt.date | None = None  # дата документа (есть в русском листе)


def parse_packing_list(rows):
    head, cells = _find_header(rows, [MODEL_WORDS, "qty/ctn"])
    if head is None:
        if _find_header(rows, [NAME_WORDS, ("упаковка", "g.w", "брутто")], without=PRICE_WORDS)[0] is not None:
            return parse_packing_list_ru(rows)
        raise ParseError("Не нашёл заголовков упаковочного листа завода (matched models, QTY/CTN).")
    c_model = _col(cells, *MODEL_WORDS)
    c_mark = _col(cells, "item on the carton", "mark", "唛头")  # маркировка коробки (в листах с «Item No.» её нет)
    if c_mark == c_model:
        c_mark = None
    c_color = _col(cells, "color")
    c_qty = _col(cells, "qty/ctn")
    c_ctn = _col(cells, "cartons", exclude=("size", "no."))
    c_nw = _col(cells, "n.w/ctn")
    c_gw = _col(cells, "g.w/ctn")
    c_total = _col(cells, "total qty")
    c_tgw = _col(cells, "total g.w")
    c_tnw = _col(cells, "total n.w")

    pl = PackingList(title=" ".join(_text(c) for c in rows[0] if _text(c)) if rows else "")
    if m := re.search(r"/\s*([A-Z]{2}\d{5,})", pl.title):
        pl.invoice = m.group(1)

    def get(row, col):
        return row[col] if col is not None and col < len(row) else None

    carton_specs = defaultdict(lambda: defaultdict(Counter))  # name -> шт./коробку -> {(брутто, нетто): коробок}
    # Смешанная коробка (образцы): одна строка с маркировкой, весом и общим количеством T, остальные —
    # только модель и штуки. В Excel эта строка первая; в PDF объединённые ячейки печатаются посередине
    # группы — поэтому строки без коробок «ждут» (pending) и состав группы добирается до суммы T.
    group, need, pending = None, 0, []

    def add_mixed(model, n):
        item = pl.items.setdefault(model, PackingItem(name=model))
        item.qty += n
        item.in_mixed += n

    for row in rows[head + 1:]:
        model = _text(get(row, c_model))
        if not model:
            total_ctn, total_qty = _num(get(row, c_ctn)), _num(get(row, c_total))
            if total_ctn and total_qty and pl.total_cartons is None:  # строка итогов
                pl.total_cartons, pl.total_qty = int(total_ctn), int(total_qty)
                pl.total_gross, pl.total_net = _num(get(row, c_tgw)), _num(get(row, c_tnw))
            continue
        qty, cartons, total = _num(get(row, c_qty)), _num(get(row, c_ctn)), _num(get(row, c_total))
        if qty is None and cartons is None and total is None:  # заголовок, повторённый на новой странице PDF
            continue
        item = pl.items.setdefault(model, PackingItem(name=model))
        if _text(get(row, c_color)):
            item.color = item.color or _text(get(row, c_color))
        if cartons is None:  # строка смешанной коробки без своих данных
            n = int(qty or 0)
            if group is not None and need > 0:
                group[1].append((model, n))
                add_mixed(model, n)
                need -= n
            else:
                pending.append((model, n))
            continue
        if qty and total is not None and total != qty * cartons:  # строка с данными смешанной коробки
            members = pending + [(model, int(qty))]
            pending = []
            group = (_text(get(row, c_mark)), members)
            pl.mixed_cartons.append(group)
            for m, n in members:
                add_mixed(m, n)
            need = int(total) - sum(n for _, n in members)
            continue
        group, need = None, 0
        if not qty:
            continue
        item.mark = item.mark or _text(get(row, c_mark))
        item.qty += int(qty * cartons)
        carton_specs[model][int(qty)][(_num(get(row, c_gw)), _num(get(row, c_nw)))] += int(cartons)
    for model, n in pending:  # без коробки — штуки всё равно учитываем (в смешанной коробке)
        add_mixed(model, n)
    for name, sizes in carton_specs.items():
        result = []
        for q, weights in sorted(sizes.items()):
            # один размер коробки с разным весом: вес — по большинству коробок
            (gw, nw), _ = weights.most_common(1)[0]
            result.append(PackingCarton(q, sum(weights.values()), gw, nw))
        pl.items[name].cartons = result
    if not pl.items:
        raise ParseError("В упаковочном листе не нашлось строк с изделиями.")
    return pl


def parse_packing_list_ru(rows):
    """Русский упаковочный лист завода: по изделию — всего штук, коробок и общий вес.

    Раскладки по коробкам в нём нет: размер коробки берём, только если штуки делятся на коробки
    поровну (иначе коробки — из китайского упаковочного листа). Строки без коробок после строки
    с коробкой — одна смешанная коробка (образцы), как в китайском листе.
    """
    head, cells = _find_header(rows, [NAME_WORDS, ("упаковка", "g.w", "брутто")], without=PRICE_WORDS)
    c_name = _col(cells, *NAME_WORDS)
    c_qty = _col(cells, *QTY_WORDS)
    c_ctn = _col(cells, "упаковка", "коробок", "мест", "cartons", "ctns")
    c_gw = _col(cells, "g.w", "брутто")
    c_nw = _col(cells, "n.w", "нетто")
    if c_qty is None:
        raise ParseError("В упаковочном листе не нашлось колонки с количеством.")

    def get(row, col):
        return row[col] if col is not None and col < len(row) else None

    header_text = " ".join(_text(c) for row in rows[:head] for c in row)
    pl = PackingList(title=" ".join(_text(c) for c in rows[head - 1] if _text(c)) if head else "", source="ru")
    if m := re.search(r"\b([A-Z]{2}\d{6,})\b", header_text):
        pl.invoice = m.group(1)
    if m := re.search(r"Дата:?\s*(\d{4}-\d{2}-\d{2}|\d{1,2}\.\d{1,2}\.\d{4}|\d{4}/\d{1,2}/\d{1,2}|[а-яё]+\.?\s*\d{1,2}[.,\s]+\d{4})",
                      header_text, re.IGNORECASE):
        pl.date = _parse_date(m.group(1))
    group, last_name = None, ""
    for row in rows[head + 1:]:
        name = _text(get(row, c_name))
        qty, cartons = _num(get(row, c_qty)), _num(get(row, c_ctn))
        first = " ".join(_text(c) for c in row[:max(c_name or 0, 1) + 1]).lower()
        if first.startswith(("общий итог", "итого", "total", "всего")):
            if qty and pl.total_qty is None:
                pl.total_qty, pl.total_cartons = int(qty), int(cartons) if cartons else None
                pl.total_gross, pl.total_net = _num(get(row, c_gw)), _num(get(row, c_nw))
            continue
        if name.lower().startswith(("общий", "итого")):
            continue
        if not name and qty is not None and last_name:
            name = last_name  # строка без названия — та же позиция в другой коробке (остаток: «75 шт, 1 коробка»)
        if not name or qty is None:
            continue
        last_name = name
        item = pl.items.setdefault(name, PackingItem(name=name))
        item.qty += int(qty)
        if cartons is None and group is not None:  # в той же коробке, что строка выше
            item.in_mixed += int(qty)
            group[1].append((name, int(qty)))
            if len(group[1]) == 2:  # первая строка группы тоже оказалась в смешанной коробке
                first = pl.items[group[1][0][0]]
                first.in_mixed += group[1][0][1]
                first.cartons = [c for c in first.cartons if c is not group[2]]
                pl.mixed_cartons.append(group)
            continue
        group = None
        if cartons:
            gw, nw = _num(get(row, c_gw)), _num(get(row, c_nw))
            per = Decimal(int(qty)) / cartons
            spec = None
            if per == int(per):
                spec = PackingCarton(int(per), int(cartons), (gw / cartons).quantize(Decimal("0.01")) if gw else None,
                                     (nw / cartons).quantize(Decimal("0.01")) if nw else None)
                item.cartons.append(spec)
            group = ("", [(name, int(qty))], spec)
    pl.mixed_cartons = [(mark, members) for mark, members, _ in pl.mixed_cartons]
    if not pl.items:
        raise ParseError("В упаковочном листе не нашлось строк с изделиями.")
    return pl


def parse_document(rows):
    """Что за документ: ("invoice", Invoice) или ("packing", PackingList). Имя файла не важно."""
    if _find_header(rows, [MODEL_WORDS, "qty/ctn"])[0] is not None:
        return "packing", parse_packing_list(rows)
    if _find_header(rows, [NAME_WORDS, QTY_WORDS, PRICE_WORDS], without=("qty/ctn",))[0] is not None:
        return "invoice", parse_invoice(rows)
    if _find_header(rows, [NAME_WORDS, ("упаковка", "g.w", "брутто")], without=PRICE_WORDS)[0] is not None:
        return "packing", parse_packing_list_ru(rows)
    raise ParseError("Не похоже ни на инвойс (нет колонок «Наименование», «Кол-во», «Цена»), "
                     "ни на упаковочный лист завода (нет «matched models», «QTY/CTN» или «Упаковка», «G.W»).")
