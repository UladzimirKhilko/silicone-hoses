"""Переносит историю из «Учет патрубков МТЗ.xlsx» (docs/11): приходы → партии, отгрузки → отгрузки МТЗ.

    python manage.py import_mtz_excel "/путь/Учет патрубков МТЗ.xlsx"

Партии отгрузок в Excel не указаны — выбираем их тем же правилом, что и сервис (stock.services.pick_batches),
из партий, пришедших к дате отгрузки.
Разовая команда: если партии уже есть, ничего не делает (флаг --force удаляет историю и грузит заново).
"""

import datetime as dt
from collections import defaultdict
from decimal import Decimal

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from catalog.models import Customer, Product, find_product
from stock.models import Adjustment, Batch, ReceiptLine, Shipment, ShipmentLine
from stock.services import pick_batches

# Столбцы приходов и что о партии известно из других документов (docs/10, docs/11).
RECEIPTS = {
    "D": (dt.date(2024, 3, 4), {"notes": "Образцы. Перенесено из «Учет патрубков МТЗ.xlsx»."}),
    "E": (dt.date(2025, 2, 8), {"invoice": "TS2407202", "factory_orders": "Q571, Q595, Q604",
                                "notes": "Перенесено из Excel (только изделия МТЗ). Образцы ГСМ/ММЗ из этой поставки — внести отдельно."}),
    "F": (dt.date(2025, 9, 18), {"number": "02/25", "manufactured": "07.2025", "passport_number": "09/002/25",
                                 "notes": "Номер партии, дата изготовления и паспорт определены по паспортам и биркам ноября 2025."}),
    "G": (dt.date(2025, 11, 4), {"notes": "Перенесено из Excel."}),
    "H": (dt.date(2025, 11, 19), {"notes": "Перенесено из Excel."}),
}
# Столбцы отгрузок: дата из шапки (месяц — строка 2, число — строка 3).
SHIPMENTS = {
    "M": dt.date(2025, 7, 8), "O": dt.date(2025, 8, 5), "P": dt.date(2025, 8, 12), "Q": dt.date(2025, 8, 26),
    "S": dt.date(2025, 9, 2), "T": dt.date(2025, 9, 8), "U": dt.date(2025, 9, 24), "W": dt.date(2025, 11, 11),
    "X": dt.date(2025, 11, 19), "Z": dt.date(2025, 12, 17), "AA": dt.date(2025, 12, 19), "AC": dt.date(2026, 3, 6),
    "AD": dt.date(2026, 3, 10), "AE": dt.date(2026, 3, 19),
}
TTN = {dt.date(2025, 11, 11): "5134413", dt.date(2025, 11, 18): "5134443", dt.date(2025, 11, 19): "5134450"}
# По ТТН 5134443 D50L100-3 (150 шт.) отгружен 18.11, а в Excel записан в столбец 11.11 — верим документу.
DATE_FIX = {("BSI D50L100-3", dt.date(2025, 11, 11)): dt.date(2025, 11, 18)}


def _n(v):
    return v if isinstance(v, (int, float)) and v else 0


class Command(BaseCommand):
    help = "Перенести историю приходов и отгрузок из Excel МТЗ"

    def add_arguments(self, parser):
        parser.add_argument("xlsx")
        parser.add_argument("--force", action="store_true", help="удалить существующие партии и отгрузки")

    @transaction.atomic
    def handle(self, xlsx, force, **opts):
        if Batch.objects.exists() and not force:
            raise CommandError("Партии уже есть. Повторный перенос — только с --force (история будет перезаписана).")
        if force:
            ShipmentLine.objects.all().delete()
            Shipment.objects.all().delete()
            Adjustment.objects.all().delete()
            Batch.objects.all().delete()

        ws = openpyxl.load_workbook(xlsx, data_only=True).active
        mtz = Customer.objects.get(name="МТЗ")
        batches = {
            col: Batch.objects.create(status=Batch.Status.RECEIVED, received_date=date, **info)
            for col, (date, info) in RECEIPTS.items()
        }
        events, unknown = [], []
        for r in range(4, ws.max_row + 1):
            name = ws[f"B{r}"].value
            if not name or not str(name).strip().lower().startswith("патрубок"):
                continue
            product = find_product(name, customer=mtz)
            if product is None:
                unknown.append(str(name))
                continue
            price_cny = Decimal(str(ws[f"C{r}"].value)) if _n(ws[f"C{r}"].value) else None
            price_byn = Decimal(str(ws[f"L{r}"].value)).quantize(Decimal("0.01")) if _n(ws[f"L{r}"].value) else None
            landed = Decimal(str(ws[f"AJ{r}"].value)).quantize(Decimal("0.01")) if _n(ws[f"AJ{r}"].value) else None
            lines = []
            for col, batch in batches.items():
                qty = int(_n(ws[f"{col}{r}"].value))
                if qty:
                    lines.append(ReceiptLine.objects.create(batch=batch, product=product, qty_expected=qty, qty_received=qty, price_cny=price_cny))
                    events.append((batch.received_date, 0, product, batch, qty))
            if landed and lines:
                # В Excel приходная цена одна на изделие — относим её к последней партии.
                last = max(lines, key=lambda line: line.batch.received_date)
                last.landed_price_byn = landed
                last.save()
            for col, date in SHIPMENTS.items():
                qty = int(_n(ws[f"{col}{r}"].value))
                if qty:
                    events.append((DATE_FIX.get((product.code, date), date), 1, product, price_byn, qty))
        if unknown:
            raise CommandError("Не найдены изделия: " + "; ".join(unknown))

        # Хронологически: приходы раньше отгрузок того же дня.
        events.sort(key=lambda e: (e[0], e[1]))
        stock = defaultdict(list)  # product_id -> [[batch, qty], …] в порядке прихода
        shipments = {}
        for date, kind, product, extra, qty in events:
            if kind == 0:
                stock[product.pk].append([extra, qty])
                continue
            shipment = shipments.get(date)
            if shipment is None:
                shipment = shipments[date] = Shipment.objects.create(
                    customer=mtz, date=date, ttn_number=TTN.get(date, ""), notes="Перенесено из «Учет патрубков МТЗ.xlsx»."
                )
            slots = stock[product.pk]
            picked = pick_batches([(b, q) for b, q in slots], qty)
            if picked is None:
                raise CommandError(f"{product}: на {date:%d.%m.%Y} не хватает товара для отгрузки {qty} шт.")
            for batch, take in picked:
                ShipmentLine.objects.create(shipment=shipment, product=product, batch=batch, qty=take, price_byn=extra)
                next(slot for slot in slots if slot[0] == batch)[1] -= take

        self.stdout.write(self.style.SUCCESS(
            f"Партий: {len(batches)}, строк прихода: {ReceiptLine.objects.count()}, "
            f"отгрузок: {len(shipments)}, строк отгрузки: {ShipmentLine.objects.count()}"
        ))
