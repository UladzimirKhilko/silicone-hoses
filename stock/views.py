import datetime as dt
from decimal import Decimal, InvalidOperation
from io import BytesIO

import openpyxl
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import F, ProtectedError, Sum
from django.core.exceptions import PermissionDenied
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST

from catalog.models import Customer, Product
from catalog.views import file_response, search_products

from . import paperwork, pricing, receipt_import, services
from .importers import ParseError, parse_invoice, parse_packing_list, read_rows
from .forms import (
    AdjustmentForm,
    CartonForm,
    AdjustmentLineFormSet,
    BatchForm,
    ReceiptLineFormSet,
    ReceiveForm,
    ShipmentForm,
    ShipmentLineFormSet,
    filled,
)
from .models import Adjustment, Batch, CartonSpec, ReceiptLine, Shipment, ShipmentLine


def costs_required(view):
    """Учётные цены и прибыль — только руководителю."""
    def wrapper(request, *args, **kwargs):
        if not pricing.can_see_costs(request.user):
            raise PermissionDenied("Учётные цены и прибыль видит только руководитель.")
        return view(request, *args, **kwargs)
    wrapper.__name__, wrapper.__doc__ = view.__name__, view.__doc__
    return login_required(wrapper)


def _decimal(raw):
    raw = (raw or "").replace(",", ".").replace(" ", "").replace("\xa0", "").strip()
    if not raw:
        return None
    try:
        return Decimal(raw)
    except InvalidOperation:
        raise ValueError(f"«{raw}» — не число")


def _last_sale_prices(product_ids):
    """{product_id: цена последней отгрузки (любому заказчику)}."""
    result = {}
    rows = (ShipmentLine.objects.filter(product_id__in=product_ids, price_byn__isnull=False)
            .order_by("shipment__date", "pk").values_list("product_id", "price_byn"))
    for product_id, price in rows:
        result[product_id] = price
    return result


# --- Приходы (партии) ---

@login_required
def batch_list(request):
    # Meta.ordering не применяется к запросам с агрегатами — сортируем явно.
    batches = (
        Batch.objects.annotate(qty=Sum("lines__qty_expected"))
        .prefetch_related("lines")
        .order_by(F("received_date").desc(nulls_first=True), "-created_at")
    )
    if request.GET.get("format") == "xlsx":
        from reports import excel

        costs = pricing.can_see_costs(request.user)
        lines = ReceiptLine.objects.select_related("batch", "product__customer").order_by(
            F("batch__received_date").desc(nulls_first=True), "batch_id", "product__code")
        rows = [[l.batch.invoice, l.batch.short_label, l.batch.get_status_display(), l.batch.received_date, l.product.code,
                 l.product.yapib, str(l.product.customer or ""), l.qty_expected, l.qty_received, l.price_cny]
                + ([l.landed_price_byn] if costs else []) for l in lines]
        return excel.table("Приходы", ["Инвойс", "Партия", "Статус", "Дата приёмки", "Изделие", "ЯПИБ", "Заказчик", "По документам",
                                       "Принято", "Цена завода, CNY"] + (["Учётная цена, BYN"] if costs else []),
                           rows, f"Приходы {dt.date.today():%Y-%m-%d}.xlsx")
    return render(request, "stock/batch_list.html", {"batches": batches})


@login_required
def batch_detail(request, pk):
    batch = get_object_or_404(Batch, pk=pk)
    lines = batch.lines.select_related("product__customer")
    shipped = {
        r["product_id"]: r["q"]
        for r in ShipmentLine.objects.filter(batch=batch).values("product_id").annotate(q=Sum("qty"))
    }
    balances = services.batch_balances([line.product_id for line in lines])
    lines = list(lines)
    sale = _last_sale_prices([line.product_id for line in lines])
    for line in lines:
        line.shipped = shipped.get(line.product_id, 0)
        line.left = balances.get((line.product_id, batch.pk), 0)
        line.amount_cny = line.price_cny * line.qty if line.price_cny is not None else None
        line.sale_price = sale.get(line.product_id)
        if line.sale_price is not None and line.landed_price_byn is not None:
            line.unit_profit = line.sale_price - line.landed_price_byn
            line.margin = (line.unit_profit / line.sale_price * 100).quantize(Decimal("0.1")) if line.sale_price else None
    priced = [line.amount_cny for line in lines if line.amount_cny is not None]
    landed = [line.landed_price_byn * line.qty for line in lines if line.landed_price_byn is not None]
    return render(request, "stock/batch_detail.html", {
        "batch": batch, "lines": lines,
        "total_cny": sum(priced) if priced else None,
        "total_landed": sum(landed) if landed else None,
        "landed_missing": sum(1 for line in lines if line.landed_price_byn is None),
        "receive_form": ReceiveForm(initial={"received_date": dt.date.today()}),
        "cartons": batch.cartons.select_related("product"),
        "carton_form": CartonForm(batch=batch),
    })


@login_required
@require_POST
def batch_info(request, pk):
    """Номер партии, дата изготовления, номер паспорта — нужны для паспорта и бирок, правятся и после приёмки."""
    batch = get_object_or_404(Batch, pk=pk)
    for field in ("number", "manufactured", "passport_number"):
        setattr(batch, field, request.POST.get(field, getattr(batch, field)).strip())
    batch.save()
    messages.success(request, "Данные партии сохранены.")
    return redirect("stock:batch_detail", pk=pk)


@costs_required
@require_POST
def batch_prices(request, pk):
    """Учётные цены партии: рассчитать от курса и расходов или ввести вручную. Можно и после приёмки."""
    batch = get_object_or_404(Batch, pk=pk)
    lines = list(batch.lines.select_related("product"))
    try:
        batch.cny_rate = _decimal(request.POST.get("cny_rate"))
        batch.extra_costs_byn = _decimal(request.POST.get("extra_costs_byn"))
        manual = {line.pk: _decimal(request.POST.get(f"landed_{line.pk}")) for line in lines}
    except ValueError as e:
        messages.error(request, f"Не сохранено: {e}.")
        return redirect(f"{batch.get_absolute_url()}#prices")
    if request.POST.get("action") == "calc":
        if not batch.cny_rate:
            messages.error(request, "Для расчёта укажите курс: сколько BYN за 1 CNY.")
            return redirect(f"{batch.get_absolute_url()}#prices")
        calculated, k = pricing.calc_landed(lines, batch.cny_rate, batch.extra_costs_byn)
        manual.update(calculated)
        no_price = [line.product.code for line in lines if line.pk not in calculated]
        messages.success(request, f"Учётные цены рассчитаны: курс {batch.cny_rate}, коэффициент расходов {k:.4f}."
                         + (f" Без цены завода (не рассчитаны): {', '.join(no_price)}." if no_price else ""))
    else:
        messages.success(request, "Учётные цены сохранены.")
    with transaction.atomic():
        batch.save(update_fields=["cny_rate", "extra_costs_byn"])
        for line in lines:
            line.landed_price_byn = pricing.money(manual[line.pk])
            line.save(update_fields=["landed_price_byn"])
    return redirect(f"{batch.get_absolute_url()}#prices")


@login_required
@require_POST
def carton_add(request, pk):
    batch = get_object_or_404(Batch, pk=pk)
    form = CartonForm(request.POST, batch=batch)
    if form.is_valid():
        spec = form.save(commit=False)
        spec.batch = batch
        spec.source = spec.source or "введено вручную"
        spec.save()
        messages.success(request, f"Коробка добавлена: {spec}.")
    else:
        messages.error(request, " ".join(e for errs in form.errors.values() for e in errs))
    return redirect("stock:batch_detail", pk=pk)


@login_required
@require_POST
def carton_delete(request, pk, carton_pk):
    get_object_or_404(CartonSpec, pk=carton_pk, batch_id=pk).delete()
    messages.success(request, "Коробка удалена.")
    return redirect("stock:batch_detail", pk=pk)


@login_required
def batch_edit(request, pk=None):
    batch = get_object_or_404(Batch, pk=pk) if pk else Batch()
    if batch.is_received and request.method == "POST":
        messages.error(request, "Партия уже принята — состав менять нельзя, только приёмку.")
        return redirect("stock:batch_detail", pk=batch.pk)
    initial = [{"product": line.product, "qty_expected": line.qty_expected, "price_cny": line.price_cny} for line in batch.lines.all()] if pk else []
    form = BatchForm(request.POST or None, instance=batch)
    formset = ReceiptLineFormSet(request.POST or None, initial=initial, prefix="lines")
    if request.method == "POST" and form.is_valid() and formset.is_valid():
        rows = filled(formset)
        products = [r["product"].pk for r in rows]
        if len(products) != len(set(products)):
            messages.error(request, "Одно изделие указано дважды — объедините строки.")
        elif not rows:
            messages.error(request, "Добавьте хотя бы одно изделие.")
        else:
            with transaction.atomic():
                batch = form.save()
                batch.lines.all().delete()
                for r in rows:
                    ReceiptLine.objects.create(batch=batch, product=r["product"], qty_expected=r["qty_expected"], price_cny=r["price_cny"])
            messages.success(request, f"Сохранено: {batch}. Товар в пути — примите его, когда придёт на склад.")
            return redirect("stock:batch_detail", pk=batch.pk)
    return render(request, "stock/batch_form.html", {"form": form, "formset": formset, "batch": batch if pk else None})


# --- Поставка из файлов завода (этап 4) ---

IMPORT_SESSION_KEY = "receipt_import_draft"


@login_required
def batch_import(request):
    """Шаг 1: загрузить инвойс и/или упаковочный лист завода."""
    if request.method == "POST":
        invoice_file, pl_file = request.FILES.get("invoice"), request.FILES.get("packing_list")
        if not invoice_file and not pl_file:
            messages.error(request, "Выберите инвойс и/или упаковочный лист завода.")
            return redirect("stock:batch_import")
        try:
            invoice = parse_invoice(read_rows(invoice_file, invoice_file.name)) if invoice_file else None
            packing = parse_packing_list(read_rows(pl_file, pl_file.name)) if pl_file else None
        except ParseError as e:
            messages.error(request, str(e))
            return redirect("stock:batch_import")
        except Exception:  # битый или чужой файл — не падаем с 500, а объясняем
            messages.error(request, "Не удалось прочитать файл. Нужен Excel (.xlsx или .xls) — инвойс или упаковочный лист завода.")
            return redirect("stock:batch_import")
        draft = receipt_import.build_draft(invoice, packing)
        draft["files"] = [f.name for f in (invoice_file, pl_file) if f]
        request.session[IMPORT_SESSION_KEY] = draft
        return redirect("stock:batch_import_preview")
    return render(request, "stock/batch_import.html")


@login_required
def batch_import_preview(request):
    """Шаг 2: проверить сопоставление и количества, выбрать изделия для неопознанных строк, создать поставку."""
    draft = request.session.get(IMPORT_SESSION_KEY)
    if not draft:
        return redirect("stock:batch_import")
    rows = receipt_import.rows_from(draft)
    header = {
        "invoice": draft["header"].get("invoice", ""),
        "factory_orders": "", "number": "", "manufactured": "", "passport_number": "",
        "notes": "Загружено из файлов: " + ", ".join(draft.get("files", [])),
    }
    if request.method == "POST":
        header = {k: request.POST.get(k, v).strip() for k, v in header.items()}
        choices = {r.key: int(request.POST[f"product_{r.key}"]) for r in rows
                   if not r.product_id and request.POST.get(f"product_{r.key}", "").isdigit()}
        missing = [r for r in rows if not r.product_id and r.key not in choices]
        if missing:
            messages.error(request, f"Выберите изделие для строк: {', '.join(r.invoice_name or r.pl_name for r in missing)}.")
        else:
            target = Batch.objects.filter(pk=request.POST.get("target") or 0, status=Batch.Status.IN_TRANSIT).first()
            batch = receipt_import.create_batch(draft, choices, header, remember=bool(request.POST.get("remember")),
                                                user=request.user, target=target)
            del request.session[IMPORT_SESSION_KEY]
            messages.success(request, f"{'Заказ обновлён по инвойсу' if target else 'Поставка создана'}: {batch}. "
                             "Товар «в пути» — примите его, когда придёт на склад.")
            return redirect("stock:batch_detail", pk=batch.pk)
    products = Product.objects.in_bulk([r.product_id for r in rows if r.product_id] + [c for r in rows for c in r.candidate_ids])
    for r in rows:
        r.product = products.get(r.product_id)
        r.candidates = [products[c] for c in r.candidate_ids if c in products]
    return render(request, "stock/batch_import_preview.html", {
        "draft": draft, "h": draft["header"], "rows": rows, "header": header,
        "warnings": receipt_import.draft_warnings(draft, rows),
        "all_products": Product.objects.select_related("customer").order_by("customer__name", "code"),
        "unmatched": sum(1 for r in rows if not r.product_id),
        "total_qty": sum(r.qty or 0 for r in rows),
        "orders": Batch.objects.filter(status=Batch.Status.IN_TRANSIT, invoice="").annotate(qty=Sum("lines__qty_expected")).order_by("-created_at"),
    })


@login_required
@require_POST
def batch_receive(request, pk):
    """Приёмка: фактическое количество по строкам; расхождение остаётся видно в партии."""
    batch = get_object_or_404(Batch, pk=pk)
    form = ReceiveForm(request.POST)
    if not form.is_valid():
        messages.error(request, "Укажите дату приёмки.")
        return redirect("stock:batch_detail", pk=pk)
    lines = list(batch.lines.select_related("product"))
    shipped = {
        r["product_id"]: r["q"] for r in ShipmentLine.objects.filter(batch=batch).values("product_id").annotate(q=Sum("qty"))
    }
    errors = []
    for line in lines:
        raw = request.POST.get(f"received_{line.pk}", "").strip()
        try:
            qty = int(raw) if raw else line.qty_expected
        except ValueError:
            errors.append(f"{line.product}: «{raw}» — не число.")
            continue
        if qty < 0:
            errors.append(f"{line.product}: количество не может быть отрицательным.")
        elif qty < shipped.get(line.product_id, 0):
            errors.append(f"{line.product}: из партии уже отгружено {shipped[line.product_id]} шт.")
        line.qty_received = qty
    if errors:
        for e in errors:
            messages.error(request, e)
        return redirect("stock:batch_detail", pk=pk)
    with transaction.atomic():
        for line in lines:
            line.save()
        batch.status = Batch.Status.RECEIVED
        batch.received_date = form.cleaned_data["received_date"]
        batch.save()
    short = [f"{line.product.code}: −{line.shortage}" for line in lines if line.shortage > 0]
    messages.success(request, "Партия принята." + (" Недостача: " + ", ".join(short) if short else ""))
    return redirect("stock:batch_detail", pk=pk)


@login_required
@require_POST
def batch_delete(request, pk):
    batch = get_object_or_404(Batch, pk=pk)
    try:
        batch.delete()
    except ProtectedError:
        messages.error(request, "Из партии уже были отгрузки или списания — удалить нельзя.")
        return redirect("stock:batch_detail", pk=pk)
    messages.success(request, "Партия удалена.")
    return redirect("stock:batch_list")


# --- Отгрузки ---

def _fill_prices(customer, rows):
    """Пустая цена = цена из последней отгрузки этого изделия этому заказчику."""
    for r in rows:
        if r.get("price_byn") is None:
            last = (
                ShipmentLine.objects.filter(shipment__customer=customer, product=r["product"], price_byn__isnull=False)
                .order_by("-shipment__date", "-pk")
                .first()
            )
            r["price_byn"] = last.price_byn if last else None
    return rows


@login_required
def shipment_list(request):
    shipments = Shipment.objects.select_related("customer").order_by("-date", "-created_at")
    if request.GET.get("customer"):
        shipments = shipments.filter(customer_id=request.GET["customer"])
    if request.GET.get("month"):
        try:
            y, m = map(int, request.GET["month"].split("-"))
        except ValueError:
            y, m = dt.date.today().year, dt.date.today().month
        shipments = shipments.filter(date__year=y, date__month=m)
    if request.GET.get("format") == "xlsx":
        from reports import excel

        costs = pricing.can_see_costs(request.user)
        lines = ShipmentLine.objects.filter(shipment__in=shipments).select_related(
            "shipment__customer", "shipment__request", "product", "batch", "spec_line__spec").order_by("shipment__date", "shipment_id", "product__code")
        lines = pricing.with_costs(lines) if costs else list(lines)
        rows = [[l.shipment.date, l.shipment.customer.name, l.shipment.ttn_number, l.product.code, l.product.yapib, l.batch.short_label,
                 l.qty, l.price_byn, l.amount, f"№ {l.spec_line.spec.number}" if l.spec_line else "", str(l.shipment.request or "")]
                + ([l.unit_cost, l.profit] if costs else []) for l in lines]
        return excel.table("Отгрузки", ["Дата", "Заказчик", "ТТН", "Изделие", "ЯПИБ", "Партия", "Шт", "Цена без НДС", "Сумма",
                                        "Спецификация", "Заявка"] + (["Учётная цена", "Прибыль"] if costs else []),
                           rows, f"Отгрузки {dt.date.today():%Y-%m-%d}.xlsx")
    shipments = list(shipments)
    lines = ShipmentLine.objects.filter(shipment__in=shipments)
    lines = pricing.with_costs(lines) if pricing.can_see_costs(request.user) else list(lines)
    by_shipment = {g.key: g.totals for g in pricing.group_by(lines, lambda line: line.shipment_id, str)}
    for shipment in shipments:
        shipment.t = by_shipment.get(shipment.pk, pricing.Totals())
    return render(request, "stock/shipment_list.html", {
        "shipments": shipments, "customers": Customer.objects.all(), "filters": request.GET,
        "total": pricing.totals(lines) if pricing.can_see_costs(request.user) else None,
        "total_qty": sum(line.qty for line in lines),
        "total_amount": sum((line.amount or 0) for line in lines),
    })


def _paperwork_warnings(shipment, lines, boxes):
    warnings = []
    for batch in {line.batch for line in lines}:
        missing = [label for label, value in (("номер партии", batch.number), ("номер паспорта", batch.passport_number),
                                              ("дату изготовления", batch.manufactured)) if not value]
        if missing:
            warnings.append(f"В партии «{batch}» не указаны: {', '.join(missing)} — заполните в карточке партии.")
    for box in boxes:
        if box.gross_kg is None:
            warnings.append(f"{box.product.code} (партия «{box.batch.short_label}»): нет данных о коробке — масса на бирке будет пустой. Добавьте коробку в карточке партии.")
    if not shipment.ttn_number:
        warnings.append("Не указан номер ТТН — в паспорте поле останется пустым.")
    if not shipment.customer.full_name:
        warnings.append(f"У заказчика «{shipment.customer}» нет полного наименования — в паспорт пойдёт краткое.")
    return warnings


@login_required
def shipment_detail(request, pk):
    shipment = get_object_or_404(Shipment.objects.select_related("customer"), pk=pk)
    lines = list(shipment.lines.select_related("product", "batch", "spec_line__spec"))
    if pricing.can_see_costs(request.user):
        lines = pricing.with_costs(lines)
    total = sum((line.amount or 0) for line in lines)
    boxes = paperwork.box_plan(sorted(lines, key=lambda line: (line.product.code, line.batch_id)))
    return render(request, "stock/shipment_detail.html", {
        "shipment": shipment, "lines": lines, "total": total, "boxes": boxes,
        "totals": pricing.totals(lines) if pricing.can_see_costs(request.user) else None,
        "box_total": sum(b.count for b in boxes),
        "paperwork_warnings": _paperwork_warnings(shipment, lines, boxes),
    })


DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@login_required
def shipment_passport(request, pk):
    shipment = get_object_or_404(Shipment.objects.select_related("customer"), pk=pk)
    name = f"Паспорт {shipment.customer} {shipment.date:%d.%m.%Y}" + (f" ТТН {shipment.ttn_number}" if shipment.ttn_number else "")
    return file_response(BytesIO(paperwork.passport_docx(shipment)), f"{name}.docx", content_type=DOCX)


@login_required
def shipment_labels(request, pk):
    shipment = get_object_or_404(Shipment.objects.select_related("customer"), pk=pk)
    name = f"Бирки {shipment.customer} {shipment.date:%d.%m.%Y}"
    return file_response(BytesIO(paperwork.labels_docx(shipment)), f"{name}.docx", content_type=DOCX)


@login_required
def shipment_create(request):
    from sales import services as sales

    initial, lines_initial = {"date": dt.date.today()}, []
    req = None
    if request.GET.get("request", "").isdigit():  # отгрузка по заявке: невыполненный остаток
        from sales.models import CustomerRequest

        req = get_object_or_404(CustomerRequest, pk=request.GET["request"])
        initial.update(customer=req.customer, request=req)
        lines_initial = [{"product": line.product, "qty": min(line.left, line.free) or line.left, "price_byn": line.price_byn}
                         for line in sales.request_rows(req) if line.left]
    form = ShipmentForm(request.POST or None, request.FILES or None, initial=initial)
    formset = ShipmentLineFormSet(request.POST or None, prefix="lines", initial=lines_initial)
    if request.method == "POST" and form.is_valid() and formset.is_valid():
        rows = filled(formset)
        if not rows:
            messages.error(request, "Добавьте хотя бы одно изделие.")
        else:
            try:
                with transaction.atomic():
                    shipment = form.save(commit=False)
                    shipment.created_by = request.user
                    shipment.save()
                    warnings = sales.apply_specs(shipment.customer, shipment.date, rows, form.cleaned_data.get("specification"))
                    services.post_shipment(shipment, _fill_prices(shipment.customer, rows))
            except services.InsufficientStock as e:
                messages.error(request, f"Не хватает товара. {e}")
            else:
                messages.success(request, f"Отгрузка проведена: {shipment}.")
                for w in warnings:
                    messages.warning(request, w)
                return redirect("stock:shipment_detail", pk=shipment.pk)
    return render(request, "stock/shipment_form.html", {"form": form, "formset": formset, "req": req})


@login_required
@require_POST
def shipment_attach_ttn(request, pk):
    shipment = get_object_or_404(Shipment, pk=pk)
    if request.FILES.get("ttn_scan"):
        if shipment.ttn_scan:
            shipment.ttn_scan.delete(save=False)
        shipment.ttn_scan = request.FILES["ttn_scan"]
    shipment.ttn_number = request.POST.get("ttn_number", shipment.ttn_number).strip()
    shipment.save()
    messages.success(request, "ТТН сохранена.")
    return redirect("stock:shipment_detail", pk=pk)


@login_required
def shipment_ttn_scan(request, pk):
    shipment = get_object_or_404(Shipment, pk=pk)
    if not shipment.ttn_scan or not shipment.ttn_scan.storage.exists(shipment.ttn_scan.name):
        raise Http404("Скан ТТН не загружен")
    ext = shipment.ttn_scan.name.rsplit(".", 1)[-1]
    return file_response(shipment.ttn_scan.open("rb"), f"ТТН {shipment.ttn_number or shipment.pk}.{ext}", as_attachment=False)


@login_required
@require_POST
def shipment_delete(request, pk):
    shipment = get_object_or_404(Shipment, pk=pk)
    shipment.delete()
    messages.success(request, f"Отгрузка удалена, товар вернулся в остаток: {shipment}.")
    return redirect("stock:shipment_list")


# --- Списания и корректировки ---

@login_required
def adjustment_list(request):
    items = Adjustment.objects.select_related("customer").annotate(qty=Sum("lines__qty")).order_by("-date", "-created_at")
    if request.GET.get("format") == "xlsx":
        from reports import excel

        from .models import AdjustmentLine

        lines = AdjustmentLine.objects.select_related("adjustment__customer", "product", "batch").order_by("-adjustment__date")
        rows = [[l.adjustment.date, l.adjustment.get_kind_display(), str(l.adjustment.customer or ""), l.product.code, l.batch.short_label,
                 l.adjustment.sign * l.qty, l.adjustment.notes] for l in lines]
        return excel.table("Списания", ["Дата", "Причина", "Заказчик", "Изделие", "Партия", "Шт (+/−)", "Комментарий"],
                           rows, f"Списания {dt.date.today():%Y-%m-%d}.xlsx")
    return render(request, "stock/adjustment_list.html", {"items": items})


@login_required
def adjustment_detail(request, pk):
    adjustment = get_object_or_404(Adjustment, pk=pk)
    return render(request, "stock/adjustment_detail.html", {"adjustment": adjustment, "lines": adjustment.lines.select_related("product", "batch")})


@login_required
def adjustment_create(request):
    form = AdjustmentForm(request.POST or None, initial={"date": dt.date.today(), "kind": request.GET.get("kind")})
    formset = AdjustmentLineFormSet(request.POST or None, prefix="lines")
    if request.method == "POST" and form.is_valid() and formset.is_valid():
        rows = filled(formset)
        if not rows:
            messages.error(request, "Добавьте хотя бы одно изделие.")
        else:
            try:
                with transaction.atomic():
                    adjustment = form.save(commit=False)
                    adjustment.created_by = request.user
                    adjustment.save()
                    services.post_adjustment(adjustment, rows)
            except services.InsufficientStock as e:
                messages.error(request, f"Не хватает товара. {e}")
            else:
                messages.success(request, f"Проведено: {adjustment}.")
                return redirect("stock:adjustment_detail", pk=adjustment.pk)
    return render(request, "stock/adjustment_form.html", {"form": form, "formset": formset})


@login_required
@require_POST
def adjustment_delete(request, pk):
    adjustment = get_object_or_404(Adjustment, pk=pk)
    if adjustment.sign > 0:
        # Удаление прихода по инвентаризации не должно увести партию в минус.
        balances = services.batch_balances([line.product_id for line in adjustment.lines.all()])
        for line in adjustment.lines.all():
            if balances.get((line.product_id, line.batch_id), 0) < line.qty:
                messages.error(request, f"{line.product}: этот товар уже отгружен — удалить нельзя.")
                return redirect("stock:adjustment_detail", pk=pk)
    adjustment.delete()
    messages.success(request, "Удалено.")
    return redirect("stock:adjustment_list")


# --- Выгрузка остатков ---

@login_required
def stock_export(request):
    products = list(search_products(request.GET))
    ids = [p.pk for p in products]
    balances = services.batch_balances(ids)
    have, transit = services.on_hand(ids, balances), services.in_transit(ids)
    costs = pricing.can_see_costs(request.user)
    value = pricing.stock_value(balances)[0] if costs else {}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Остатки"
    from sales.services import sale_prices

    prices = sale_prices((p.pk, p.customer_id) for p in products)
    ws.append(["Код", "ЯПИБ", "Заказчик", "Тип", "Размеры", "Цвет", "Остаток, шт", "Цена без НДС, BYN", "Сумма, BYN", "Откуда цена",
               "В пути, шт"] + (["Сумма по учётным ценам, BYN"] if costs else []))
    for p in products:
        price, source = prices[p.pk, p.customer_id]
        qty = have.get(p.pk, 0)
        ws.append([p.code, p.yapib or "", p.customer.name if p.customer_id else "", p.get_kind_display(), p.size_label,
                   p.get_color_display(), qty, float(price) if price is not None else None,
                   float(price * qty) if price is not None and qty else None, source, transit.get(p.pk, 0)]
                  + ([float(pricing.money(value[p.pk])) if p.pk in value else None] if costs else []))
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=9):
        for cell in row:
            cell.number_format = "#,##0.00"
    for col, width in zip("ABCDEFGHIJKL", (28, 10, 14, 16, 26, 10, 12, 14, 14, 18, 12, 16)):
        ws.column_dimensions[col].width = width
    return _xlsx_response(wb, f"ostatki-{dt.date.today():%Y-%m-%d}.xlsx")


def _xlsx_response(wb, filename):
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


# --- Отчёт «Продажи и прибыль» ---

MONTHS = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]


def _month_label(d):
    return f"{MONTHS[d.month - 1]} {d.year}"


def _report_period(params):
    today = dt.date.today()
    default_from = dt.date(today.year - (1 if today.month < 12 else 0), today.month % 12 + 1, 1)  # 12 месяцев
    def parse(name, default):
        try:
            return dt.date.fromisoformat(params.get(name, ""))
        except ValueError:
            return default
    return parse("date_from", default_from), parse("date_to", today)


@costs_required
def profit_report(request):
    date_from, date_to = _report_period(request.GET)
    lines = ShipmentLine.objects.filter(shipment__date__range=(date_from, date_to)).select_related("shipment__customer", "product")
    customer = request.GET.get("customer", "")
    if customer.isdigit():
        lines = lines.filter(shipment__customer_id=customer)
    lines = pricing.with_costs(lines.order_by("shipment__date", "pk"))
    by_month = pricing.group_by(lines, lambda line: line.shipment.date.replace(day=1), lambda line: _month_label(line.shipment.date))
    by_customer = sorted(pricing.group_by(lines, lambda line: line.shipment.customer_id, lambda line: line.shipment.customer.name),
                         key=lambda g: -g.totals.amount)
    by_product = sorted(pricing.group_by(lines, lambda line: line.product_id, lambda line: line.product.code),
                        key=lambda g: -g.totals.amount)
    if request.GET.get("format") == "xlsx":
        return _profit_xlsx(date_from, date_to, by_month, by_customer, by_product, lines)
    return render(request, "stock/profit_report.html", {
        "date_from": date_from, "date_to": date_to, "customer": customer, "customers": Customer.objects.all(),
        "total": pricing.totals(lines), "by_month": by_month, "by_customer": by_customer, "by_product": by_product,
    })


def _profit_xlsx(date_from, date_to, by_month, by_customer, by_product, lines):
    from openpyxl.styles import Font

    head = ["Шт", "Сумма без НДС, BYN", "Себестоимость, BYN", "Прибыль, BYN", "Маржа, %", "Без учётной цены, шт", "Без цены продажи, шт"]

    def row(t):
        return [t.qty, float(t.amount), float(pricing.money(t.cost)), float(pricing.money(t.profit)),
                float(t.margin) if t.margin is not None else None, t.qty_no_cost or None, t.qty_no_price or None]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, first, groups in (("По месяцам", "Месяц", by_month), ("По заказчикам", "Заказчик", by_customer), ("По изделиям", "Изделие", by_product)):
        ws = wb.create_sheet(title)
        ws.append([f"Продажи и прибыль {date_from:%d.%m.%Y} — {date_to:%d.%m.%Y}"])
        ws["A1"].font = Font(bold=True)
        ws.append([first] + head)
        for g in groups:
            ws.append([g.label] + row(g.totals))
        ws.append(["Итого"] + row(pricing.totals(lines)))
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        ws.column_dimensions["A"].width = 30
        for col in "BCDEFGH":
            ws.column_dimensions[col].width = 16
    ws = wb.create_sheet("Строки отгрузок")
    ws.append(["Дата", "Заказчик", "ТТН", "Изделие", "Партия", "Шт", "Цена без НДС", "Сумма", "Учётная цена", "≈", "Себестоимость", "Прибыль"])
    for line in lines:
        ws.append([line.shipment.date, line.shipment.customer.name, line.shipment.ttn_number, line.product.code, line.batch.short_label,
                   line.qty, float(line.price_byn) if line.price_byn is not None else None,
                   float(line.amount) if line.amount is not None else None,
                   float(line.unit_cost) if line.unit_cost is not None else None, "≈" if line.cost_approx else "",
                   float(line.cost) if line.cost is not None else None, float(pricing.money(line.profit)) if line.profit is not None else None])
        ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY"
    for col, width in zip("ABCDEFGHIJKL", (11, 14, 10, 26, 12, 7, 12, 12, 12, 4, 13, 12)):
        ws.column_dimensions[col].width = width
    return _xlsx_response(wb, f"pribyl-{date_from:%Y-%m-%d}-{date_to:%Y-%m-%d}.xlsx")
