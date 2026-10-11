"""Выгрузка таблиц в Excel одним вызовом: заголовок, жирная шапка, ширина колонок, формат чисел."""

import datetime as dt
from decimal import Decimal
from io import BytesIO
from urllib.parse import quote

import openpyxl
from django.http import HttpResponse
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="E8EEF8")


def _plain(value):
    if isinstance(value, Decimal):
        return float(value)
    return value


def add_sheet(wb, title, headers, rows, caption=None, widths=None, total_row=False):
    ws = wb.create_sheet(title[:31])
    r = 1
    if caption:
        ws.cell(1, 1, caption).font = Font(bold=True, size=12)
        r = 3
    for c, h in enumerate(headers, 1):
        cell = ws.cell(r, c, h)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = ws.cell(r + 1, 1)
    for row in rows:
        r += 1
        for c, value in enumerate(row, 1):
            cell = ws.cell(r, c, _plain(value))
            if isinstance(value, (dt.date, dt.datetime)):
                cell.number_format = "DD.MM.YYYY"
            elif isinstance(value, Decimal) or isinstance(value, float):
                cell.number_format = "#,##0.00"
    if total_row and rows:
        for cell in ws[r]:
            cell.font = Font(bold=True)
    for c, h in enumerate(headers, 1):
        longest = max((len(str(row[c - 1] if row[c - 1] is not None else "")) for row in rows[:500]), default=0)
        width = (widths or {}).get(c) or min(48, max(10, len(str(h)) + 2, longest + 2))
        ws.column_dimensions[get_column_letter(c)].width = width
    return ws


def response(wb, filename):
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        wb.remove(wb["Sheet"])
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    ascii_name = filename.encode("ascii", "ignore").decode() or "export.xlsx"
    resp["Content-Disposition"] = f"attachment; filename=\"{ascii_name}\"; filename*=utf-8''{quote(filename)}"
    return resp


def table(title, headers, rows, filename, caption=None, widths=None):
    wb = openpyxl.Workbook()
    add_sheet(wb, title, headers, rows, caption, widths)
    return response(wb, filename)


def wants_xlsx(request):
    return request.GET.get("format") == "xlsx"
