"""Сводка, отчёты, выгрузки в Excel, новый интерфейс (этап 6)."""

import datetime as dt
import io
from decimal import Decimal

import openpyxl
import pytest
from django.contrib.auth.models import Group
from django.core.management import call_command
from django.urls import reverse

from catalog.models import Customer, Product
from stock import services as stock
from stock.models import Batch, ReceiptLine, Shipment

from . import analytics as an


@pytest.fixture
def mtz(db):
    return Customer.objects.create(name="МТЗ")


@pytest.fixture
def product(mtz):
    return Product.objects.create(code="BSI D75L100-2", yapib="24.1119", customer=mtz)


@pytest.fixture
def client_in(client, django_user_model):
    client.force_login(django_user_model.objects.create_user("m", password="pass12345", first_name="Ольга"))
    return client


def received(product, qty, date=dt.date(2025, 1, 1)):
    batch = Batch.objects.create(status=Batch.Status.RECEIVED, received_date=date)
    ReceiptLine.objects.create(batch=batch, product=product, qty_expected=qty, qty_received=qty)
    return batch


def ship(customer, product, qty, date, price="5.00"):
    s = Shipment.objects.create(customer=customer, date=date)
    stock.post_shipment(s, [{"product": product, "qty": qty, "price_byn": Decimal(price)}])
    return s


def test_monthly_pivot(mtz, product):
    amk = Customer.objects.create(name="Амкадор")
    received(product, 1000)
    ship(mtz, product, 100, dt.date(2025, 3, 5))
    ship(mtz, product, 50, dt.date(2025, 3, 20))
    ship(amk, product, 10, dt.date(2025, 5, 1), price="8")
    months, rows, totals, grand = an.monthly_shipments(dt.date(2025, 3, 1), dt.date(2025, 5, 31))
    assert months == [dt.date(2025, 3, 1), dt.date(2025, 4, 1), dt.date(2025, 5, 1)]
    assert [r.label for r in rows] == ["МТЗ", "Амкадор"]  # по сумме
    assert [(c.qty, c.amount) for c in rows[0].cell_list(months)] == [(150, Decimal("750.00")), (0, 0), (0, 0)]
    assert (grand.qty, grand.amount) == (160, Decimal("830.00"))
    _, by_product, _, _ = an.monthly_shipments(dt.date(2025, 3, 1), dt.date(2025, 5, 31), by="product")
    assert by_product[0].label == "BSI D75L100-2"


def test_low_stock_by_average_and_minimum(mtz, product):
    other = Product.objects.create(code="BSI D60L80-2", customer=mtz, min_stock=500)
    quiet = Product.objects.create(code="BSI D50L100-3", customer=mtz)
    received(product, 1000)
    received(other, 300)
    received(quiet, 50)
    today = dt.date(2025, 9, 15)
    for m in range(3, 9):  # по 120 шт в месяц полгода
        ship(mtz, product, 120, dt.date(2025, m, 10))
    rows = {r["product"]: r for r in an.low_stock(today)}
    # осталось 1000 − 720 = 280 → хватит на 280 / 120 ≈ 2,3 мес < 4 → заказать до 6 мес: 720 − 280 = 440
    assert rows[product]["months"] == Decimal("2.3") and rows[product]["suggested"] == 440
    assert rows[other]["months"] is None and rows[other]["suggested"] == 200  # ниже минимума 500
    assert quiet not in rows  # продаж нет, минимума нет


def test_bar_chart_geometry():
    points = [(dt.date(2025, m, 1), Decimal(v), "") for m, v in ((1, 0), (2, 1000), (3, 2500))]
    chart = an.bar_chart(points)
    assert [g["label"] for g in chart["grid"]] == ["0", "1 тыс", "2 тыс", "3 тыс"]
    bars = chart["bars"]
    assert bars[0]["path"] == "" and bars[2]["show_value"] and not bars[1]["show_value"]
    assert bars[2]["cx"] - bars[1]["cx"] > 24  # столбцы не слипаются


def test_dashboard_and_reports_pages(client_in, mtz, product):
    received(product, 100)
    ship(mtz, product, 10, dt.date.today())
    html = client_in.get(reverse("dashboard")).content.decode()
    assert "Добрый день, Ольга" in html and "На складе" in html and 'class="chart"' in html
    assert "Прибыль" not in html  # менеджеру прибыль не показываем
    for name in ("reports:index", "reports:shipments_monthly", "reports:low_stock", "help"):
        assert client_in.get(reverse(name)).status_code == 200


def test_navigation_marks_active_and_hides_profit_for_manager(client_in):
    html = client_in.get(reverse("catalog:product_list")).content.decode()
    assert 'class="nav-item active"' in html and "Остатки" in html
    assert reverse("stock:profit_report") not in html


def test_boss_sees_profit_in_nav_and_dashboard(client, django_user_model, mtz, product):
    call_command("setup_roles")
    boss = django_user_model.objects.create_user("b", password="pass12345")
    boss.groups.add(Group.objects.get(name="Руководитель"))
    client.force_login(boss)
    html = client.get(reverse("dashboard")).content.decode()
    assert reverse("stock:profit_report") in html and "Руководитель" in html


def _sheet(resp):
    assert resp["Content-Type"].startswith("application/vnd.openxmlformats")
    return openpyxl.load_workbook(io.BytesIO(resp.content)).active


@pytest.mark.parametrize("url", ["stock:shipment_list", "stock:batch_list", "stock:adjustment_list", "sales:request_list", "reports:low_stock"])
def test_excel_everywhere(client_in, mtz, product, url):
    received(product, 100)
    ship(mtz, product, 10, dt.date(2025, 3, 1))
    ws = _sheet(client_in.get(reverse(url), {"format": "xlsx"}))
    assert ws.max_column >= 5


def test_excel_contents(client_in, mtz, product):
    received(product, 100)
    ship(mtz, product, 10, dt.date(2025, 3, 1))
    ws = _sheet(client_in.get(reverse("stock:shipment_list"), {"format": "xlsx"}))
    assert [c.value for c in ws[2]][:7] == [dt.datetime(2025, 3, 1), "МТЗ", None, "BSI D75L100-2", "24.1119", "01.01.2025", 10]
    assert "Учётная цена" not in [c.value for c in ws[1]]  # менеджеру — без себестоимости
    wb = openpyxl.load_workbook(io.BytesIO(client_in.get(reverse("reports:shipments_monthly"),
                                                         {"date_from": "2025-01-01", "date_to": "2025-03-31", "format": "xlsx"}).content))
    assert wb.sheetnames == ["Сумма без НДС", "Количество"]
    ws = wb["Количество"]
    assert [c.value for c in ws[4]] == ["МТЗ", 0, 0, 10, 10]
    ws = _sheet(client_in.get(product.get_absolute_url(), {"format": "xlsx"}))
    assert ws.cell(4, 4).value == 100 and ws.cell(5, 4).value == -10


def test_login_page_renders(client, db):
    html = client.get(reverse("login")).content.decode()
    assert "Вход в систему" in html and "login-form" in html
